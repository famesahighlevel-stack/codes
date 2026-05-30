import requests
import time
import json
import re
import pandas as pd
import threading
import tkinter as tk
from tkinter import messagebox
import customtkinter as cctk
from tkcalendar import Calendar
from datetime import timedelta, datetime, timezone, time as dt_time
import pytz
from openpyxl.styles import PatternFill, Font, Alignment
from concurrent.futures import ThreadPoolExecutor, as_completed
from html import unescape
import os

# ---------------------------
# CONFIG: GHL Cuentas
# ---------------------------
ACCOUNTS = [
    {
        "name": "R2.1",
        "location_id": "gPO6FiXFjyul3qojXQKh",
        "stage_id": "dea2dacb-97d4-4165-b95b-5a02f6a87155",
        "custom_field": "fydz3DjQN2pib2axnFZe",
        "dataventa_id": "7veOYT8o670WKCIbRd92",
        "token": "pit-aa32ae57-f021-4345-9c17-599788ca222d",
        "secuencia_cf": "CDP7RyHYQIhVRNtmuYOX",
        "anuncio_cf": "EsjAWJ6IbHmCB47SbGTe",
        "primer_mensaje_cf": "ILBHzfqHvQpVXZqOqPTt"
    },
    {
        "name": "R2.2",
        "location_id": "HK2q1x0usZEQbnuJCD6B",
        "stage_id": "a5b50462-e82e-4256-b83d-d613ad20abcd",
        "custom_field": "U6OdqsQDVQTUDxxNBtvp",
        "dataventa_id": "1L7km8XXq3V1kDMRyvCC",
        "token": "pit-c8f89986-1a34-4178-9064-6f678a697a01",
        "secuencia_cf": "deDXbK9Nw5khnuk679CM",
        "anuncio_cf": "vzIYG6B7mypHxXM9kLGi",
        "primer_mensaje_cf": "gMCJJq0vMZYN0cz1BnUj"
    },
    {
        "name": "R2.3",
        "location_id": "jXN4id73HVqpa75YOR1N",
        "stage_id": "d94817e9-a7fb-4ea3-bed0-d1f117001825",
        "custom_field": "xftnXlHb41aDvIx8N26d",
        "dataventa_id": "JJYIQaDMKpMMekORmcqy",
        "token": "pit-cfd6cb24-6fee-441d-b01f-415160eb6f7b",
        "secuencia_cf": "ENXxDMYOkFK7XnjzrdOU",
        "anuncio_cf": "r8sOhHm5PKNtXv65jSaN",
        "primer_mensaje_cf": "vphhYZkFfAiaSrgfTiA3"
    },
    {
        "name": "R1.1",
        "location_id": "GmvsWG2a09UJjYwzFwN7",
        "stage_id": "c8d3a128-f03d-4bce-b61d-50593b7c4ebc",
        "custom_field": "k6xXzN1ksd16i1n68o4P",
        "dataventa_id": "kHdjdcIduLvTr3nL6DgZ",
        "token": "pit-cc6a4560-665e-4063-b12e-7dd3ec412570",
        "secuencia_cf": "rG3qADV2DReag3vE4LnZ",
        "anuncio_cf": "hN5VqUs4cgWP4tNEdnTI",
        "primer_mensaje_cf": "at9WccT1xJ4tJRm5KpbK"
    },
    {
        "name": "R1.2",
        "location_id": "9rHHeTsNpfJuiUkOoLdM",
        "stage_id": "59f6eff0-f06b-4f5a-b2ae-21f50ec8af32",
        "custom_field": "o7giXoy1LK8KMuzH2FNi",
        "dataventa_id": "DeNNFP4LihWoLaIpG0B2",
        "token": "pit-51105ced-165a-437d-bf76-37c9ec75f00e",
        "secuencia_cf": "AAYTXtJX7jRHPn0VDqVH",
        "anuncio_cf": "IELF1xRsnl1nWvoHBRHY",
        "primer_mensaje_cf": "uEKcNGLJvv7znfVJb2Z4"
    },
    {
        "name": "R1.3",
        "location_id": "riT0De9iiwhd84gSRco3",
        "stage_id": "3b346e40-01ea-416e-98ff-60ee147aded1",
        "custom_field": "f310g4Z4A1OxHl3KYLRv",
        "dataventa_id": "4tDKaHgodZXQ4ewzjJyh",
        "token": "pit-58f93f09-e7a8-40d0-8ce6-32bf2f2b87b4",
        "secuencia_cf": "7s52aRnEuz8T3v0H3IlN",
        "anuncio_cf": "Vki3QtqsNJtYcRC4Fkzt",
        "primer_mensaje_cf": "78eH2yy4INV88QguossT"
    }
]

VENDEDOR_MAP = {
    "MARIA RENE SANTA CRUZ COSAJAY": "MARIA SANTACRUZ",
    "HENRY ESTUARDO PACHECO ARIANO": "HENRY PACHECO",
    "ROSA LIDIA PEREZ": "ROSA PEREZ",
    "CLEMENCIA ROCIO MIZA": "CLEMENCIA MIZA",
    "GRECIA SARAI FLORES FLORES": "GRECIA FLORES",
    "EVELYN PAOLA DARDON MORATAYA": "EVELYN DARDON",
    "MARVIN FRANCISCO CARRERA PINEDA": "MARVIN CARRERA",
    "YEIMI NOHEMI HERNANDEZ GOMEZ": "YEIMI HERNANDEZ",
    "ANA KARINA VELASQUEZ CALDERON DE CANAS": "ANA VELASQUEZ",
    "JHONATAN ISRAEL BARRIOS MUÑOZ": "JHONATAN BARRIOS",
    "CRISTIAN OMAR SANTOS ROSALES": "Cristian Santos",
    "LOURDES CAROLINA CAMPOS REYES": "Lourdes Campos",
    "CARLOS ARMANDO GIL TAJIN": "Carlos Gil",
    "OSCAR DANIEL IXCAYAU AGUILAR": "Oscar Ixcayau",
    "SIN ASIGNAR": "SinAsignar",
    "BYRON SILVERIO ORTIZ MENDOZA": "BYRON ORTIZ",
    "ESTHER DEL CARMEN LOPEZ CRUZ": "ESTHER LOPEZ",
    "SONIA MARIBEL CHIROY": "SONIA CHIROY",
    "MARVIN DAVID CANEL HERNANDEZ": "MARVIN CANEL",
    "BLANCA  RUTILIA YACAB AC": "BLANCA YACAB",
    "PEDRO JOSUE YAX VILLATORO": "PEDRO VILLATORO",
    "VICTOR MIGUEL ANGEL ASTURIAS TZAMOL": "VICTOR ANGEL",
    "ERICKA VICTORIA GONZALEZ LÓPEZ": "ERICKA GONZALEZ",
    "KAREN YULISA MENCOS RODAS": "KAREN MENCOS",
    "FABIANA MENDOZA LOPEZ": "FABIANA LOPEZ",
    "STEPHANIE DENNES AJUCHAN HERNANDEZ": "Stephanie Ajuchan",
    "ALISON MELISA AJANEL JOLON": "Alyson Ajanel",
    "IRMA LOURDES DE PAZ": "Lourdes Paz",
    "DENIS SEBASTIAN AJUCHAN SANTOS": "Denis Ajuchan",
    "SANDRA PATRICIA GUTIERREZ BLANCO": "Sandra Gutierrez",
    "FELIX  ANTONIO ACEITUNO BARRIENTOS": "FELIX ACEITUNO",
    "JUAN LUIS COTZAJAY TIJE": "JUAN COTZAJAY",
    "BRIAN SALVADOR CANO RAMIREZ": "BRIAN CANO",
    "JAMNIA GALILEA MORALES ORELLANA": "GALILEA MORALES",
    "WILMER ALEXANDER CAJAS JUAREZ": "WILMER CAJAS",
    "ADA LUZ MARIA MENDEZ SOSA": "Ada luz Mendez",
    "CHRISTIAN OLIVER GONZALEZ ORDOÑEZ": "Christian Gonzalez",
    "JESSICA ILEANA GARCIA PALACIOS": "JESSICA GARCIA",
    "GERBER ROMUALDO YAX VILLATORO": "GERBER YAX",
    "HAZEL BARINIA AGUILAR GONZALEZ": "HAZEL AGUILAR",
    "OSCAR JOSUE HERNANDEZ PEREZ": "Oscar Hernandez",
    "WALTER NEHEMIAS GUERRA": "WALTER GUERRA",
    "YESSICA ALEJANDRA CARRERA PINEDA": "YESSICA CARRERA",
    "HIDELKY AJIN": "HidelkyAjin",
    "BYRON ALEXIS MAYOR": "BYRON MAYOR",
    "YARELIN BARRAZA ARIAS": "YARELIN ARIAS",
    "YOSELIN EUFEMIA BARRAZA ARIAS": "YOSELIN BARRAZA",
    "YENDY MIREYA CUMAR CASTRO": "YENDY CUMAR",
    "ANGEL DANIEL LOPEZ PATZAN": "Angel Lopez",
    "BRYAN ARMANDO LOPEZ PATZAN": "Bryan Lopez",
    "HILLARY GISSELL RUCAL GALLINA": "Hillary Rucal",
    "KIMBERLY JESSENIA REGUAN RABAY": "Kimberly Reguán",
    "LUIS DAVID QUEXEL GIL": "Luis Quexel",
    "WENDY ARELIS ALQUIJAY GIL": "Wendy Alquijay",
    "CARLOS GIL TAJIN": "Carlos Gil",
    "ROSALINDA EUSEBIA RAMIREZ TEZEN": "Rosalinda Ramirez",
    "JULIO ALEJANDRO AJXUP GIL": "Julio Ajxup",
    "JOCARI ANASOL LOPEZ SICAL": "Jocari Lopez",
    "ODILIA NINETTE CALEL CARAU": "ODILIA NINETH CALEL",
}

API_VERSION_OPPS = "2023-02-21"
API_VERSION_CONTACTS = "2021-07-28"
GUATEMALA_TZ = pytz.timezone("America/Guatemala")

# ---------------------------
# CONFIG: Facebook Ads
# ---------------------------
FB_ACCESS_TOKEN = "EAAQlsSqsOJkBQ8ELEZCLxm0CPEiqaUSoYw0oHB7ML7xvufZBn2B6t1bizlxBtv8gjc1r4bHiqlV0AHtbI9FDLTRhivFwpDr2xMzk7Waj8htSHanBW63gZCOPwPoVOgZBuNurP2ZB6vegJGxYRAmIR2Wp2JbPZAn9u4CVdHZC06TzaIngYeZAh3n4iaMtw3SyZA7XsNAZDZD"
FB_API_VERSION = "v19.0"
FB_AD_ACCOUNTS = ["act_622689460111355", "act_934171589566820"]
FB_USD_ACCOUNTS = ["act_934171589566820"]
USD_TO_GTQ = 7.8
FB_BASE_URL = f"https://graph.facebook.com/{FB_API_VERSION}"
FB_FIELDS_INSIGHTS = "ad_id,ad_name,adset_id,adset_name,campaign_name,impressions,spend,clicks,actions,date_start,date_stop"

# Regex para extracción
ANUNCIO_REGEX = re.compile(r"([A-Z]\d{3,4}[A-Z]\d{3})", re.IGNORECASE)
SECUENCIA_REGEX = re.compile(r"([A-Z]\d\.\d)", re.IGNORECASE)

# ---------------------------
# Backend Functions: General
# ---------------------------
def clean_html(raw_html):
    if not raw_html:
        return ""
    return unescape(re.sub(r'<[^>]+>', '', raw_html)).strip()

def format_date_ghl(val):
    if not val:
        return ""
    if isinstance(val, (int, float)):
        try:
            return datetime.fromtimestamp(val / 1000, tz=timezone.utc).strftime("%d/%m/%Y")
        except:
            return str(val)
    s = str(val).strip()
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return f"{s[8:10]}/{s[5:7]}/{s[0:4]}"
    return s

def get_yyyy_mm_dd(val):
    if not val:
        return ""
    if isinstance(val, (int, float)):
        try:
            return datetime.fromtimestamp(val / 1000, tz=timezone.utc).strftime("%Y-%m-%d")
        except:
            return ""
    s = str(val).strip()
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    return ""

def calculate_nit(nit, tel1, ghl_phone):
    n, t1, gp = str(nit).strip().upper(), str(tel1).strip(), str(ghl_phone).strip()
    res = n
    if n in ("CF", "C/F", ""):
        res = t1
    if not res:
        if len(gp) >= 12:
            res = gp[4:12]
        elif len(gp) > 4:
            res = gp[4:]
    return res

def get_custom_value(field):
    if not field or not isinstance(field, dict):
        return ""
    if "fieldValueDate" in field and field["fieldValueDate"]:
        return format_date_ghl(field["fieldValueDate"])
    if "fieldValueString" in field and field["fieldValueString"]:
        return str(field["fieldValueString"])
    if "fieldValue" in field and field["fieldValue"] is not None:
        v = field["fieldValue"]
        if isinstance(v, (int, float)) and v > 1000000000000:
            return format_date_ghl(v)
        return str(v)
    if "value" in field and field["value"] is not None:
        v = field["value"]
        if isinstance(v, list):
            return ", ".join(map(str, v))
        return str(v)
    return ""

def get_custom_fields_map(location_id, token):
    url = f"https://services.leadconnectorhq.com/locations/{location_id}/customFields?model=opportunity"
    headers = {"Authorization": f"Bearer {token}", "Version": API_VERSION_OPPS, "Accept": "application/json"}
    try:
        r = requests.get(url, headers=headers, timeout=30)
        if r.status_code != 200:
            return {}
        return {f.get("id"): f.get("name") for f in r.json().get("customFields", []) if isinstance(f, dict)}
    except:
        return {}

def get_users_by_location(location_id, token, version=API_VERSION_OPPS):
    url = f"https://services.leadconnectorhq.com/users/?locationId={location_id}"
    headers = {"Authorization": f"Bearer {token}", "Version": version, "Accept": "application/json"}
    try:
        r = requests.get(url, headers=headers, timeout=30)
        if r.status_code != 200:
            return {}
        users = r.json().get("users", [])
        user_map = {}
        for u in users:
            uid = u.get("id")
            first = u.get("firstName", "") or ""
            last = u.get("lastName", "") or ""
            name = f"{first} {last}".strip() or u.get("email", "Desconocido")
            user_map[uid] = name
        return user_map
    except:
        return {}

def safe_post(url, token, payload, version):
    headers = {"Authorization": f"Bearer {token}", "Version": version, "Content-Type": "application/json"}
    for attempt in range(1, 6):
        try:
            r = requests.post(url, headers=headers, json=payload, timeout=30)
            if r.status_code in (200, 201):
                return r.json()
            if r.status_code == 429:
                time.sleep(attempt * 2)
                continue
            return {"__error_status": r.status_code, "__error_text": r.text}
        except:
            time.sleep(attempt * 1.5)
            continue
    return {}

def parse_dataventa(dv_str):
    try:
        data = json.loads(dv_str)
        nit, dep, mun = data.get("nit", ""), data.get("departamento", ""), data.get("municipio", "")
        t1, t2, fv, nom = data.get("tel1", ""), data.get("tel2", ""), data.get("fechaVenta", ""), data.get("nombre", "")
        p_cols = {}
        all_slots = ["Camas y Combos SKU", "Cantidad Camas y Combo SKU", "Camas y Combos SKU1", "Cantidad Camas y Combo SKU1", "Cocinas SKU", "Cantidad Cocinas SKU", "Cocinas SKU1", "Cantidad Cocinas SKU1", "Salas SKU", "Cantidad Salas SKU", "Salas SKU1", "Cantidad Salas SKU1"]
        for i, p in enumerate(data.get("productos", [])):
            if i * 2 < len(all_slots):
                p_cols[all_slots[i*2]] = p.get("sku", "")
                p_cols[all_slots[i*2 + 1]] = p.get("cantidad", "")
            else:
                p_cols[f"Producto_{i+1} SKU"] = p.get("sku", "")
                p_cols[f"Cantidad Producto_{i+1}"] = p.get("cantidad", "")
        return nit, dep, mun, t1, t2, fv, nom, p_cols
    except:
        return "", "", "", "", "", "", "", {}

def get_mapped_vendedor(raw_vendedor):
    if not raw_vendedor:
        return "SinAsignar"
    key = str(raw_vendedor).strip().upper()
    return VENDEDOR_MAP.get(key, raw_vendedor)

def parse_ventas_unnested(dv_str, contact_id, opp_id, ghl_phone, vendedor, ghl_name="", sale_date_str=""):
    data = {}
    if dv_str:
        try:
            data = json.loads(dv_str)
        except:
            pass
    nit_j, t1 = data.get("nit", ""), data.get("tel1", "")
    try:
        total_docto = float(str(data.get("Total_General", 0)).replace(',', ''))
    except:
        total_docto = 0.0
    vendedor_final = get_mapped_vendedor(vendedor)
    base_metadata = {
        "ID CONTACTO": contact_id,
        "NIT": calculate_nit(nit_j, t1, ghl_phone),
        "NOMBRE": data.get("nombre", ghl_name),
        "TEL1": t1,
        "TEL2": data.get("tel2", ""),
        "VENDEDOR": vendedor_final,
        "MUNICIPIO": data.get("municipio", ""),
        "DIRECCION": data.get("direccion", ""),
        "RCF": "000000000000",
        "canal": data.get("canal", ""),
        "DEPARTAMENTO": data.get("departamento", ""),
        "FECHA": format_date_ghl(data.get("fechaVenta", sale_date_str)),
        "MARCA": data.get("marca", ""),
        "UBICACION": data.get("ubicacion", ""),
        "ANILLO": data.get("anillo", ""),
        "COMENTARIOS": data.get("COMENTARIO", ""),
        "ID Oportunidad": opp_id,
        "BODEGAF": str(data.get("bodega", ""))[:4] if data.get("bodega") else "",
        "TOTAL DOCTO": total_docto
    }
    empty_metadata = {k: "" for k in base_metadata.keys()}
    prods = data.get("productos", [])
    if not prods:
        return [base_metadata]
    rows = []
    for i, p in enumerate(prods):
        row = dict(base_metadata) if i == 0 else dict(empty_metadata)
        sku_f = str(p.get("sku", ""))
        parts = sku_f.split("_")
        try:
            precio_combo = float(str(p.get("precio", 0)).replace(',', ''))
        except:
            precio_combo = 0.0
        row.update({
            "SKU": parts[0] if parts else sku_f,
            "DESCRIPCION": parts[1] if len(parts) > 1 else "",
            "Cantidad de combo": p.get("cantidad", ""),
            "PRECIO COMBO": precio_combo
        })
        rows.append(row)
    return rows

def make_utc_range(start_date, end_date):
    start_local = GUATEMALA_TZ.localize(datetime.combine(start_date, dt_time.min))
    end_local = GUATEMALA_TZ.localize(datetime.combine(end_date, dt_time.max))
    return start_local.astimezone(pytz.UTC).isoformat().replace("+00:00", "Z"), end_local.astimezone(pytz.UTC).isoformat().replace("+00:00", "Z")

# ---------------------------
# Backend Functions: GHL
# ---------------------------
def extraer_datos_anuncio(text):
    if not text:
        return "", ""
    match_anu = ANUNCIO_REGEX.search(text)
    if not match_anu:
        return "", ""
    anuncio = match_anu.group(1)
    tipo_post = ""
    pattern_tipo = re.escape(anuncio) + r"\.([^.]+)"
    match_tipo = re.search(pattern_tipo, text, re.IGNORECASE)
    if match_tipo:
        tipo_post = match_tipo.group(1)
    return anuncio, tipo_post

def extraer_secuencia(text):
    if not text:
        return ""
    match_sec = SECUENCIA_REGEX.search(text.upper())
    return match_sec.group(1) if match_sec else ""

def fetch_contacts_for_account(acc, start_utc, end_utc, log_callback):
    token, loc, acc_name = acc["token"], acc["location_id"], acc["name"]
    sec_cf, anu_cf, pm_cf = acc["secuencia_cf"], acc["anuncio_cf"], acc["primer_mensaje_cf"]
    log_callback(f"Extraer Contactos: {acc_name}...")
    u_map = get_users_by_location(loc, token, version=API_VERSION_OPPS)
    all_contacts, page, limit = [], 1, 100
    url = "https://services.leadconnectorhq.com/contacts/search"
    while True:
        payload = {
            "locationId": loc,
            "page": page,
            "pageLimit": limit,
            "filters": [
                {"field": "dateAdded", "operator": "range", "value": {"gt": start_utc, "lt": end_utc}}
            ]
        }
        res = safe_post(url, token, payload, API_VERSION_CONTACTS)
        if not res or (isinstance(res, dict) and res.get("__error_status")):
            break
        contacts = res.get("contacts", [])
        if not isinstance(contacts, list) or not contacts:
            break
        all_contacts.extend(contacts)
        page += 1
        if len(contacts) < limit:
            break

    formatted_contacts = []
    for c in all_contacts:
        uid = c.get("assignedTo")
        assigned_name = u_map.get(uid, "") if uid else ""
        date_iso, date_fmt = c.get("dateAdded"), ""
        if date_iso:
            dt_local = datetime.fromisoformat(date_iso.replace("Z", "+00:00")).astimezone(GUATEMALA_TZ)
            date_fmt = f"{dt_local.day}/{dt_local.month:02d}/{dt_local.year}"
        secuencia_raw, anuncio_raw, primer_mensaje_texto = "", "", ""
        for cf in c.get("customFields", []):
            cid, val = cf.get("id"), get_custom_value(cf)
            if cid == sec_cf:
                secuencia_raw = val.strip()
            elif cid == anu_cf:
                anuncio_raw = val.strip()
            elif cid == pm_cf:
                primer_mensaje_texto = val.strip()
        anuncio, tipo_post = extraer_datos_anuncio(anuncio_raw)
        if not anuncio and primer_mensaje_texto:
            anuncio, tipo_post = extraer_datos_anuncio(primer_mensaje_texto)
        secuencia = extraer_secuencia(secuencia_raw)
        formatted_contacts.append({
            "id": c.get("id", ""),
            "dateAdded": date_fmt,
            "assignedToName": assigned_name,
            "secuencia": secuencia,
            "Anuncio": anuncio,
            "tipo_post": tipo_post
        })
    return formatted_contacts

def fetch_for_account(acc, ghl_start, ghl_end, client_start, client_end, log_callback):
    token, loc, stage, cfield, dv_id, acc_name = acc["token"], acc["location_id"], acc["stage_id"], acc["custom_field"], acc["dataventa_id"], acc["name"]
    log_callback(f"Extraer Ventas: {acc_name}...")
    u_map = get_users_by_location(loc, token, version=API_VERSION_OPPS)
    cf_names = get_custom_fields_map(loc, token)
    all_opps, page, limit = [], 1, 100
    url = "https://services.leadconnectorhq.com/opportunities/search"
    while True:
        payload = {
            "locationId": loc,
            "page": page,
            "limit": limit,
            "filters": [
                {"group": "AND", "filters": [
                    {"field": "pipeline_stage_id", "operator": "eq", "value": stage},
                    {"field": "status", "operator": "eq", "value": "won"},
                    {"field": f"custom_fields.{cfield}", "operator": "range", "value": {"gte": ghl_start, "lte": ghl_end}}
                ]}
            ],
            "sort": [{"field": "date_added", "direction": "desc"}],
            "additionalDetails": {"notes": True}
        }
        res = safe_post(url, token, payload, API_VERSION_OPPS)
        if not res or (isinstance(res, dict) and res.get("__error_status")):
            break
        opps = res.get("opportunities", [])
        if not isinstance(opps, list):
            break
        all_opps.extend(opps)
        page += 1
        if len(opps) < limit:
            break
    r_opps, r_ventas = [], []
    for op in all_opps:
        if not isinstance(op, dict):
            continue
        try:
            opp_cfs = op.get("customFields") or op.get("custom_fields") or []
            sale_date_iso, sale_date_str, dv_str, cf_data = "", "", "", {}
            for cf in opp_cfs:
                fid, val = cf.get("id"), get_custom_value(cf)
                if fid == cfield:
                    sale_date_str = val
                    sale_date_iso = get_yyyy_mm_dd(cf.get("fieldValueDate") or cf.get("fieldValue") or cf.get("fieldValueString"))
                if fid == dv_id:
                    dv_str = str(cf.get("fieldValue") or cf.get("fieldValueString") or "")
                fname = cf_names.get(fid, fid)
                if fname and fname.strip().lower() != "id de oportunidad":
                    cf_data[fname] = val
            if not sale_date_iso or not (client_start <= sale_date_iso <= client_end):
                continue

            assigned_id = op.get("assignedTo") or op.get("assigned_to") or op.get("assigned_to_id")
            vendedor_raw = u_map.get(assigned_id) or assigned_id or ""

            gnam = op.get("contact", {}).get("name", "") if isinstance(op.get("contact"), dict) else ""
            opp_id_val = op.get("id", "")
            dv_data = json.loads(dv_str) if dv_str else {}
            row = {
                "secuencia": acc_name,
                "fase": op.get("pipelineStageName", "Cierre de Venta"),
                "Valor del cliente potencial": op.get("monetaryValue", 0),
                "asignado": vendedor_raw,
                "Creado": format_date_ghl(op.get("createdAt")),
                "Ultimo Actualizado": format_date_ghl(op.get("updatedAt")),
                "Seguidores": "",
                "Notas": " | ".join([clean_html(n.get("body", "")) for n in op.get("notes", []) if isinstance(n, dict)]),
                "etiquetas": ", ".join(op.get("tags", [])) if isinstance(op.get("tags"), list) else "",
                "estado": op.get("status", ""),
                "ID de contacto": op.get("contactId", ""),
                "Cliente": gnam,
                "Cod": str(opp_id_val)[:10],
                "MARCA": dv_data.get("marca", ""),
                "ANILLO": dv_data.get("anillo", ""),
                "UBICACION": dv_data.get("ubicacion", ""),
                "Mes": int(sale_date_iso[5:7]) if sale_date_iso else "",
                "DataVenta": dv_str,
                "ID de oportunidad": opp_id_val
            }
            row.update(cf_data)
            nit_j, dep, mun, t1, t2, fv_j, nom_j, p_cols = parse_dataventa(dv_str)
            gp = op.get("contact", {}).get("phone", "") if isinstance(op.get("contact"), dict) else ""
            f_final = format_date_ghl(sale_date_str or fv_j)
            row.update({
                "NIT": calculate_nit(nit_j, t1, gp),
                "Departamento": dep,
                "Municipio": mun,
                "Telefono 1": t1,
                "Telefono 2": t2,
                "Fecha": f_final,
                "Fecha de Venta": f_final
            })
            if nom_j:
                row["Cliente"] = nom_j
            row.update(p_cols)
            r_opps.append(row)
            r_ventas.extend(parse_ventas_unnested(dv_str, op.get("contactId", ""), op.get("id", ""), gp, vendedor_raw, gnam, sale_date_str))
        except:
            continue
    return r_opps, r_ventas

# ---------------------------
# Backend Functions: Facebook
# ---------------------------
def extraer_precio_fb(nombre):
    if not isinstance(nombre, str):
        return None
    m = re.search(r'\.(\d+)', nombre)
    return int(m.group(1)) if m else None

def fb_api_get(url, params):
    params["access_token"] = FB_ACCESS_TOKEN
    try:
        r = requests.get(url, params=params, timeout=30)
        return r.json()
    except:
        return {}

def obtener_insights(account, fecha_desde, fecha_hasta, log_callback):
    log_callback(f"Extraer Insights FB: {account}...")
    url = f"{FB_BASE_URL}/{account}/insights"
    params = {
        "fields": FB_FIELDS_INSIGHTS,
        "level": "ad",
        "limit": 500,
        "time_range": json.dumps({"since": fecha_desde, "until": fecha_hasta}),
        "time_increment": 1,
        "filtering": '[{"field":"spend","operator":"GREATER_THAN","value":0}]'
    }
    data = []
    while True:
        js = fb_api_get(url, params)
        if "data" not in js:
            break
        data.extend(js["data"])
        if "paging" in js and "next" in js["paging"]:
            url = js["paging"]["next"]
            params = {}
        else:
            break
    return data

def obtener_presupuestos_adsets(adset_ids):
    budget_map = {}
    if not adset_ids:
        return budget_map
    for i in range(0, len(adset_ids), 50):
        block = adset_ids[i:i+50]
        batch = [{"method": "GET", "relative_url": f"{asid}?fields=daily_budget,lifetime_budget"} for asid in block]
        r = requests.post(f"{FB_BASE_URL}/", data={"access_token": FB_ACCESS_TOKEN, "batch": json.dumps(batch)})
        responses = r.json()
        if not isinstance(responses, list):
            continue
        for resp in responses:
            if resp.get("code") != 200:
                continue
            body = json.loads(resp.get("body", "{}"))
            asid = body.get("id")
            budget = float(body.get("daily_budget") or body.get("lifetime_budget") or 0) / 100.0
            if asid:
                budget_map[str(asid)] = budget
        time.sleep(0.3)
    return budget_map

def obtener_creatives(ad_ids):
    creative_map = {}
    for i in range(0, len(ad_ids), 50):
        block = ad_ids[i:i+50]
        batch = [{"method": "GET", "relative_url": f"{aid}?fields=creative"} for aid in block]
        r = requests.post(f"{FB_BASE_URL}/", data={"access_token": FB_ACCESS_TOKEN, "batch": json.dumps(batch)})
        responses = r.json()
        if not isinstance(responses, list):
            continue
        for resp in responses:
            if resp.get("code") != 200:
                continue
            body = json.loads(resp.get("body", "{}"))
            ad_id, creative_id = body.get("id"), body.get("creative", {}).get("id")
            if ad_id and creative_id:
                creative_map[str(ad_id)] = str(creative_id)
        time.sleep(0.3)
    return creative_map

def obtener_paginas(creatives):
    page_map = {}
    if not creatives:
        return page_map
    for i in range(0, len(creatives), 50):
        block = creatives[i:i+50]
        batch = [{"method": "GET", "relative_url": f"{cid}?fields=object_story_spec,actor_id,effective_object_story_id"} for cid in block]
        r = requests.post(f"{FB_BASE_URL}/", data={"access_token": FB_ACCESS_TOKEN, "batch": json.dumps(batch)})
        responses = r.json()
        if not isinstance(responses, list):
            continue
        for resp in responses:
            if resp.get("code") != 200:
                continue
            body = json.loads(resp.get("body", "{}"))
            cid, oss, actor_id, post_id = body.get("id"), body.get("object_story_spec", {}), body.get("actor_id"), body.get("effective_object_story_id")
            page_id = oss.get("page_id") or oss.get("page", {}).get("id") or oss.get("link_data", {}).get("page_id") or oss.get("video_data", {}).get("page_id")
            if not page_id and post_id and "_" in post_id:
                page_id = post_id.split("_")[0]
            final_p_id = str(page_id) if page_id else (str(actor_id) if actor_id else None)
            if final_p_id:
                page_map[str(cid)] = final_p_id
        time.sleep(0.3)
    return page_map

def obtener_nombres_paginas(page_ids):
    names, page_ids = {}, [pid for pid in page_ids if pid]
    if not page_ids:
        return names
    for i in range(0, len(page_ids), 50):
        block = page_ids[i:i+50]
        batch = [{"method": "GET", "relative_url": f"{pid}?fields=name,username,about"} for pid in block]
        r = requests.post(f"{FB_BASE_URL}/", data={"access_token": FB_ACCESS_TOKEN, "batch": json.dumps(batch)})
        responses = r.json()
        if not isinstance(responses, list):
            continue
        for resp in responses:
            if resp.get("code") != 200:
                continue
            body = json.loads(resp.get("body", "{}"))
            pid = str(body.get("id")) if body.get("id") else None
            name = body.get("name") or body.get("username") or body.get("about") or "Sin Nombre"
            if pid:
                names[pid] = name
        time.sleep(0.3)
    return names

def obtener_paginas_autorizadas():
    names, url, params = {}, f"{FB_BASE_URL}/me/accounts", {"limit": 100}
    while True:
        js = fb_api_get(url, params)
        if "data" not in js:
            break
        for item in js["data"]:
            pid, pname = str(item.get("id")), item.get("name")
            if pid and pname:
                names[pid] = pname
        if "paging" in js and "next" in js["paging"]:
            url = js["paging"]["next"]
            params = {}
        else:
            break
    return names

# ---------------------------
# DATE PICKER
# ---------------------------
class FloatingRangePicker(cctk.CTkFrame):
    def __init__(self, parent, title):
        super().__init__(parent, corner_radius=12)
        self.start_date, self.end_date, self.pop = None, None, None
        cctk.CTkLabel(self, text=title, font=("Segoe UI", 16, "bold")).pack(anchor="w", padx=15, pady=(10,2))
        self.entry_frame = cctk.CTkFrame(self, fg_color="transparent")
        self.entry_frame.pack(padx=15, pady=(2,10), fill="x")
        self.entry_start = cctk.CTkEntry(self.entry_frame, placeholder_text="Inicio", height=32, corner_radius=8, font=("Segoe UI", 11), justify="center", state="readonly")
        self.entry_start.pack(side="left", expand=True, fill="x", padx=(0,2))
        self.entry_start.bind("<Button-1>", lambda e: self.open_calendar(self.entry_start, "start"))
        cctk.CTkLabel(self.entry_frame, text="→", font=("Segoe UI", 12, "bold")).pack(side="left")
        self.entry_end = cctk.CTkEntry(self.entry_frame, placeholder_text="Fin", height=32, corner_radius=8, font=("Segoe UI", 11), justify="center", state="readonly")
        self.entry_end.pack(side="left", expand=True, fill="x", padx=(2,0))
        self.entry_end.bind("<Button-1>", lambda e: self.open_calendar(self.entry_end, "end"))

    def open_calendar(self, target_entry, date_type):
        if self.pop:
            self.close_calendar()
        self.update_idletasks()
        self.pop = tk.Toplevel(self)
        self.pop.overrideredirect(True)
        self.pop.attributes("-topmost", True)
        self.pop.geometry(f"300x320+{target_entry.winfo_rootx()}+{target_entry.winfo_rooty() + 35}")
        self.pop.grab_set()
        container = cctk.CTkFrame(self.pop, corner_radius=10, border_width=2, border_color="#76933C", fg_color="#ffffff")
        container.pack(fill="both", expand=True)
        header = tk.Frame(container, bg="#ffffff", height=30)
        header.pack(fill="x", padx=8, pady=2)
        tk.Label(header, text="ELEGIR FECHA", font=("Segoe UI", 9, "bold"), fg="#333333", bg="#ffffff").pack(side="left")
        tk.Button(header, text="✕", font=("Arial", 9), bd=0, bg="#ffffff", command=self.close_calendar).pack(side="right")
        self.cal = Calendar(container, selectmode="day", date_pattern="yyyy-mm-dd", background='white', foreground='black', selectbackground='#1890ff', selectforeground='white', borderwidth=0)
        self.cal.pack(pady=5, padx=10, fill="both", expand=True)
        self.cal.bind("<<CalendarSelected>>", lambda e: self._on_date_selected(target_entry, date_type))
        self.pop.focus_set()

    def close_calendar(self):
        if self.pop:
            self.pop.grab_release()
            self.pop.destroy()
            self.pop = None

    def _on_date_selected(self, target_entry, date_type):
        date_obj = datetime.strptime(self.cal.get_date(), "%Y-%m-%d").date()
        if date_type == "start":
            self.start_date = date_obj
        else:
            self.end_date = date_obj
        target_entry.configure(state="normal")
        target_entry.delete(0, "end")
        target_entry.insert(0, str(date_obj))
        target_entry.configure(state="readonly")
        self.after(200, self.close_calendar)

# ---------------------------
# APP
# ---------------------------
class App(cctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("DUPAZA PRO")
        self.geometry("380x700")
        cctk.set_appearance_mode("dark")
        cctk.set_default_color_theme("green")
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)
        header = cctk.CTkFrame(self, height=45, corner_radius=0)
        header.grid(row=0, column=0, sticky="ew")
        cctk.CTkLabel(header, text="📊 DUPAZA PRO", font=("Segoe UI", 18, "bold")).pack(pady=8)
        body = cctk.CTkScrollableFrame(self, fg_color="transparent")
        body.grid(row=1, column=0, sticky="nsew", padx=8, pady=8)
        body.grid_columnconfigure(0, weight=1)
        self.contacts_picker = FloatingRangePicker(body, "📇 CONTACTOS GHL")
        self.contacts_picker.grid(row=0, column=0, pady=3, sticky="ew")
        self.sales_picker = FloatingRangePicker(body, "💰 VENTAS GHL")
        self.sales_picker.grid(row=1, column=0, pady=3, sticky="ew")
        self.fb_picker = FloatingRangePicker(body, "🔵 FACEBOOK ADS")
        self.fb_picker.grid(row=2, column=0, pady=3, sticky="ew")
        self.generate_btn = cctk.CTkButton(self, text="🚀 GENERAR EXCEL", height=38, font=("Segoe UI", 14, "bold"), corner_radius=10, command=self.start_process)
        self.generate_btn.grid(row=2, column=0, pady=8, padx=15, sticky="ew")
        logs_frame = cctk.CTkFrame(self, height=80, corner_radius=10)
        logs_frame.grid(row=3, column=0, sticky="ew", padx=10, pady=(0,10))
        self.console = cctk.CTkTextbox(logs_frame, height=60, font=("Consolas", 10))
        self.console.pack(fill="both", expand=True, padx=5, pady=5)
        self.log("LISTO.")

    def log(self, txt):
        hour = datetime.now().strftime("%H:%M:%S")
        self.console.configure(state="normal")
        self.console.insert("end", f"[{hour}] {txt}\n")
        self.console.see("end")
        self.console.configure(state="disabled")

    def start_process(self):
        has_sales = self.sales_picker.start_date and self.sales_picker.end_date
        has_contacts = self.contacts_picker.start_date and self.contacts_picker.end_date
        has_fb = self.fb_picker.start_date and self.fb_picker.end_date
        if not (has_sales or has_contacts or has_fb):
            messagebox.showwarning("Atención", "Elija al menos un rango de fechas para generar el reporte.")
            return
        self.generate_btn.configure(state="disabled", text="🚀 PROCESANDO...")
        threading.Thread(target=self.execute_logic, daemon=True, args=(has_sales, has_contacts, has_fb)).start()

    def execute_logic(self, has_sales, has_contacts, has_fb):
        try:
            self.log("Iniciando extracción modular...")
            res_o, res_v, res_c, res_fb = [], [], [], []
            with ThreadPoolExecutor(max_workers=5) as ex:
                futures = []
                f_map = {}
                if has_sales:
                    sd_opp, ed_opp = self.sales_picker.start_date, self.sales_picker.end_date
                    s_iso_o, e_iso_o = sd_opp.strftime("%Y-%m-%d"), ed_opp.strftime("%Y-%m-%d")
                    ghl_s_o = sd_opp.strftime("%Y-%m-%dT00:00:00.000Z")
                    ghl_e_o = ed_opp.strftime("%Y-%m-%dT23:59:59.999Z")
                    for acc in ACCOUNTS:
                        fut = ex.submit(fetch_for_account, acc, ghl_s_o, ghl_e_o, s_iso_o, e_iso_o, self.log)
                        futures.append(fut)
                        f_map[fut] = ("opp", acc)
                if has_contacts:
                    sd_con, ed_con = self.contacts_picker.start_date, self.contacts_picker.end_date
                    s_u_c, e_u_c = make_utc_range(sd_con, ed_con)
                    for acc in ACCOUNTS:
                        fut = ex.submit(fetch_contacts_for_account, acc, s_u_c, e_u_c, self.log)
                        futures.append(fut)
                        f_map[fut] = ("con", acc)
                if has_fb:
                    fb_s, fb_h = self.fb_picker.start_date.strftime("%Y-%m-%d"), self.fb_picker.end_date.strftime("%Y-%m-%d")
                    me_pages = obtener_paginas_autorizadas()
                    for acc in FB_AD_ACCOUNTS:
                        fut = ex.submit(obtener_insights, acc, fb_s, fb_h, self.log)
                        futures.append(fut)
                        f_map[fut] = ("fb", acc)
                for f in as_completed(futures):
                    type, acc_data = f_map[f]
                    if type == "opp":
                        o, v = f.result()
                        res_o.extend(o)
                        res_v.extend(v)
                    elif type == "con":
                        res_c.extend(f.result())
                    elif type == "fb":
                        insights = f.result()
                        acc_id = acc_data
                        if insights:
                            ad_ids = list({i["ad_id"] for i in insights if "ad_id" in i})
                            adset_ids = list({i["adset_id"] for i in insights if "adset_id" in i})
                            creative_map = obtener_creatives(ad_ids)
                            budget_map = obtener_presupuestos_adsets(adset_ids)
                            p_map = obtener_paginas(list(set(creative_map.values())))
                            all_page_names = {**me_pages, **obtener_nombres_paginas(list(set(p_map.values())))}
                            is_usd = acc_id in FB_USD_ACCOUNTS
                            for ins in insights:
                                aid = str(ins.get("ad_id"))
                                asid = str(ins.get("adset_id"))
                                ad_n = ins.get("ad_name", "")
                                camp = ins.get("campaign_name", "")
                                pid = p_map.get(creative_map.get(aid))
                                pname = all_page_names.get(pid)
                                conv = next((a["value"] for a in ins.get("actions", []) if a["action_type"] == "onsite_conversion.messaging_conversation_started_7d"), 0)
                                anu, tpost = extraer_datos_anuncio(ad_n)
                                spend_orig = float(ins.get("spend", 0))
                                budget_orig = budget_map.get(asid, 0)
                                if is_usd:
                                    spend_q, budget_q = spend_orig * USD_TO_GTQ, budget_orig * USD_TO_GTQ
                                    spend_usd, budget_usd = spend_orig, budget_orig
                                else:
                                    spend_q, budget_q = spend_orig, budget_orig
                                    spend_usd, budget_usd = "", ""
                                res_fb.append({
                                    "ID del anuncio": aid,
                                    "ID de la página": pid,
                                    "Nombre de la página": pname,
                                    "Nombre de la campaña": camp,
                                    "Nombre del conjunto": ins.get("adset_name"),
                                    "Nombre del anuncio": ad_n,
                                    "Día": ins.get("date_start"),
                                    "Contactos mensajes nuevos": conv,
                                    "Importe gastado Q": spend_q,
                                    "Presupuesto Q": budget_q,
                                    "Importe gastado $": spend_usd,
                                    "Presupuesto $": budget_usd,
                                    "Inicio informe": ins.get("date_start"),
                                    "Fin informe": ins.get("date_stop"),
                                    "codigo": anu,
                                    "precio": extraer_precio_fb(ad_n),
                                    "tipo_post": tpost,
                                    "SECUENCIA": extraer_secuencia(camp)
                                })
            if res_o or res_c or res_fb:
                self.generate_excel(res_o, res_v, res_c, res_fb)
            else:
                self.log("Sin datos.")
        except Exception as e:
            self.log(f"Error: {str(e)}")
        finally:
            self.after(0, lambda: self.generate_btn.configure(state="normal", text="🚀 GENERAR EXCEL"))

    def generate_excel(self, res_o, res_v, res_c, res_fb):
        self.log("Compilando...")
        df_o, df_v, df_c, df_fb = pd.DataFrame(res_o), pd.DataFrame(res_v), pd.DataFrame(res_c), pd.DataFrame(res_fb)
        head = ["secuencia", "fase", "Valor del cliente potencial", "asignado", "Creado", "Ultimo Actualizado", "Seguidores", "Notas", "etiquetas", "estado", "Fecha de Venta", "NIT", "Camas y Combos SKU", "Cantidad Camas y Combo SKU", "Camas y Combos SKU1", "Cantidad Camas y Combo SKU1", "Cocinas SKU", "Cantidad Cocinas SKU", "Cocinas SKU1", "Cantidad Cocinas SKU1", "Salas SKU", "Cantidad Salas SKU", "Salas SKU1", "Cantidad Salas SKU1"]
        tail = ["", "Departamento", "Municipio", "Telefono 1", "Telefono 2", "ID de oportunidad", "ID de contacto", "Cliente", "Mes", "Cod", "DataVenta", "Fecha", "MARCA", "ANILLO", "UBICACION"]
        if not df_o.empty:
            if "" not in df_o.columns:
                df_o[""] = ""
            for c in head + tail:
                if c not in df_o.columns:
                    df_o[c] = ""
            extra = [c for c in df_o.columns if c not in head and c not in tail]
            df_o = df_o[head + extra + tail]
        v_cols = ["ID CONTACTO", "NIT", "NOMBRE", "TEL1", "TEL2", "VENDEDOR", "MUNICIPIO", "DIRECCION", "RCF", "canal", "DEPARTAMENTO", "FECHA", "SKU", "DESCRIPCION", "Cantidad de combo", "MARCA", "UBICACION", "ANILLO", "COMENTARIOS", "ID Oportunidad", "BODEGAF", "TOTAL DOCTO", "PRECIO COMBO"]
        if not df_v.empty:
            for c in v_cols:
                if c not in df_v.columns:
                    df_v[c] = ""
            df_v = df_v[v_cols]
        c_cols = ["id", "dateAdded", "assignedToName", "secuencia", "Anuncio", "tipo_post"]
        if not df_c.empty:
            for c in c_cols:
                if c not in df_c.columns:
                    df_c[c] = ""
            df_c = df_c[c_cols]
        fb_cols_head = ["ID del anuncio", "ID de la página", "Nombre de la página", "Nombre de la campaña", "Nombre del conjunto", "Nombre del anuncio", "Día", "Contactos mensajes nuevos", "Importe gastado Q", "Presupuesto Q", "Importe gastado $", "Presupuesto $", "Inicio informe", "Fin informe"]
        fb_cols_tail = ["codigo", "precio", "tipo_post", "SECUENCIA"]
        if not df_fb.empty:
            for c in fb_cols_head + fb_cols_tail:
                if c not in df_fb.columns:
                    df_fb[c] = ""
            df_fb = df_fb[fb_cols_head + fb_cols_tail]
        fn = f"reporte_Dupaza_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        with pd.ExcelWriter(fn, engine='openpyxl') as writer:
            if not df_o.empty:
                df_o.to_excel(writer, sheet_name='REPORTE', index=False)
            if not df_v.empty:
                df_v.to_excel(writer, sheet_name='VENTAS', index=False)
            if not df_c.empty:
                df_c.to_excel(writer, sheet_name='CONTACTOS', index=False)
            if not df_fb.empty:
                df_fb.to_excel(writer, sheet_name='FACEBOOK ADS', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='Hoja1', index=False)
            if not df_v.empty:
                ws_v = writer.book['VENTAS']
                idx_bus = len(v_cols) + 1
                ws_v.cell(row=1, column=idx_bus).value = "BUSQUEDA"
                h_f = PatternFill(start_color="76933C", end_color="76933C", fill_type="solid")
                h_font = Font(bold=True, color="FFFFFF")
                for cell in ws_v[1]:
                    cell.fill, cell.font, cell.alignment = h_f, h_font, Alignment(horizontal="center")
                bg_f = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
                cur_f = None
                for r in range(2, ws_v.max_row + 1):
                    if ws_v.cell(row=r, column=1).value:
                        ws_v.cell(row=r, column=idx_bus).value = f"=VLOOKUP(T{r},Hoja1!A:A,1,FALSE)"
                        cur_f = bg_f if cur_f is None else None
                    if cur_f:
                        for c in range(1, idx_bus + 1):
                            ws_v.cell(row=r, column=c).fill = cur_f
        self.log(f"¡EXITO! {fn}")
        messagebox.showinfo("ÉXITO", f"Excel generado:\n{fn}")

if __name__ == "__main__":
    app = App()
    app.mainloop()
